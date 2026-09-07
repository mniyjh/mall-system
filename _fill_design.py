# -*- coding: utf-8 -*-
import os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

BASE = r"D:\20260907实训\商城系统\goods"
SRC = os.path.join(BASE, "docs", "系统设计文档.docx")
OUT = os.path.join(BASE, "docs", "系统设计文档-完成版.docx")
DIAG = os.path.join(BASE, "diagrams")
DIAG_D = os.path.join(BASE, "diagrams", "design")

doc = docx.Document(SRC)
body_parent = doc.paragraphs[0]._parent
CENTER = WD_ALIGN_PARAGRAPH.CENTER

# ================= 内容 =================
SEC11 = ("本文档依据《商城系统需求分析规格说明书》，采用面向对象方法对商城系统进行详细设计，明确系统的软件架构、"
         "类图、时序图、数据库设计及界面设计，为后续的系统实现、编码与测试提供依据，是系统开发的重要技术文档。")

SEC12 = ("商城系统是一款基于浏览器/服务器（B/S）架构的网上商品交易系统。系统默认仅存在一家商户，面向管理员、"
         "运营人员、买家三类角色，围绕商品上架、购买、发货、收货评价四大核心业务流程，实现商品管理、订单管理、"
         "购物车、下单、发货、评价与售后等功能，为交易双方提供一个完整的在线商品交易平台。")

SEC2 = ("根据需求分析，系统包含管理员、运营人员、买家三类角色，主要功能需求如下：\n"
        "（1）管理员：负责用户管理、角色管理等系统级功能。\n"
        "（2）运营人员：负责商品管理（添加、修改、上架、下架商品）、订单管理、发货管理、售后管理等业务功能。\n"
        "（3）买家：负责查询商品、添加到购物车、下单、确认收货、评价商品、申请售后等功能。\n"
        "系统围绕商品上架、购买、发货、收货评价四大业务流程展开，其功能性需求与非功能性需求详见《商城系统需求分析规格说明书》。")

ARCH_TEXT = ("商城系统采用浏览器/服务器（B/S）三层架构，自上而下分为表现层、业务逻辑层、数据访问层和数据层。"
             "表现层负责与用户交互，接收用户请求并展示处理结果；业务逻辑层封装系统的核心业务逻辑；数据访问层负责对数据库的"
             "读写操作；数据层存储系统的各类数据。系统架构如图3-1所示。")

BIZ_TEXT = "系统包含商品上架、购买、发货、收货评价四大核心业务流程，各流程活动图如图3-2至图3-5所示。"

BIZ_FLOWS = [
    ("02_activity_shangjia.png", "图3-2 商品上架流程活动图"),
    ("03_activity_goumai.png", "图3-3 购买流程活动图"),
    ("04_activity_fahuo.png", "图3-4 发货流程活动图"),
    ("05_activity_shouhuo.png", "图3-5 收货评价流程活动图"),
]

CLASS_TEXT = ("通过对需求分析中的用例进行分析，抽象出系统的实体类及其关系，得到系统的类图（领域模型），如图3-6所示。"
              "系统主要包括用户、角色、商品分类、商品、购物车、订单、订单明细、评价、售后、物流等实体类。")

SEQ_TEXT = "现就添加商品、下单、发货、申请售后四个核心用例进行时序图分析，如图3-7至图3-10所示。"

SEQ_DIAGRAMS = [
    ("03_seq_addgoods.png", "图3-7 添加商品时序图"),
    ("04_seq_order.png", "图3-8 下单时序图"),
    ("05_seq_ship.png", "图3-9 发货时序图"),
    ("06_seq_aftersale.png", "图3-10 申请售后时序图"),
]

MODULE_HEADER = ["模块编号", "模块名称", "功能说明", "接口说明"]
MODULES = [
    ("M01", "用户管理", "系统用户的新增、修改、删除、查询及停用启用", "用户登录、用户增删改查"),
    ("M02", "角色管理", "系统角色及权限的维护与分配", "角色增删改查、权限分配"),
    ("M03", "商品管理", "商品的添加、修改、上架、下架与查询", "商品增删改查、上下架"),
    ("M04", "订单管理", "订单的查询、发货与状态管理", "订单查询、发货处理"),
    ("M05", "售后管理", "售后申请的审核与处理", "售后查询、售后处理"),
    ("M06", "购物车管理", "买家购物车的添加、修改与删除", "购物车增删改查"),
    ("M07", "订单模块", "买家下单、订单查询与取消", "提交订单、订单查询"),
    ("M08", "评价模块", "买家对已收货商品的评价", "提交评价、评价查询"),
]

UI_TEXT = "系统的主要页面原型如图3-11至图3-15所示，包括登录页面、商城首页、购物车页面、后台商品管理页面与订单管理页面。"

UI_PAGES = [
    ("proto_login.png", "图3-11 登录页面原型"),
    ("proto_home.png", "图3-12 商城首页原型"),
    ("proto_cart.png", "图3-13 购物车页面原型"),
    ("proto_product_admin.png", "图3-14 商品管理页面原型"),
    ("proto_order_admin.png", "图3-15 订单管理页面原型"),
]

ER_TEXT = ("系统的概念模型分为两个模型：其一为用户权限管理模型（用户、角色），如图4-1所示；其二为主要业务实体模型"
           "（商品分类、商品、购物车、订单、订单明细、评价、售后、物流），如图4-2所示。")

ER_DIAGRAMS = [
    ("07_er_permission.png", "图4-1 用户权限E-R图"),
    ("08_er_business.png", "图4-2 业务实体E-R图"),
]

LOGIC_TEXT = "系统的数据表清单如表4-1所示。"
LOGIC_HEADER = ["表名", "中文名称", "功能描述", "主要字段"]
LOGIC_TABLES = [
    ("sys_user", "用户表", "存储系统用户信息", "user_id, username, password, role_id, status"),
    ("sys_role", "角色表", "存储系统角色信息", "role_id, role_name, role_desc"),
    ("category", "商品分类表", "存储商品分类信息", "category_id, category_name, parent_id"),
    ("goods", "商品表", "存储商品信息", "goods_id, goods_name, price, stock, status"),
    ("cart", "购物车表", "存储买家购物车项", "cart_id, user_id, goods_id, quantity"),
    ("orders", "订单表", "存储订单主信息", "order_id, order_no, user_id, total_amount, status"),
    ("order_item", "订单明细表", "存储订单商品明细", "item_id, order_id, goods_id, quantity"),
    ("review", "评价表", "存储商品评价信息", "review_id, order_id, goods_id, rating"),
    ("after_sale", "售后表", "存储售后申请信息", "sale_id, order_id, goods_id, status"),
    ("shipping", "物流表", "存储订单物流信息", "shipping_id, order_id, tracking_no"),
]

PHYSICAL_HEADER = ["字段名称", "数据类型", "长度", "主键/外键", "允许为空", "中文名称"]

def cols(*fields):
    return [list(f) for f in fields]

PHYSICAL = [
    ("表4-2 用户表（sys_user）结构", cols(
        ("user_id", "bigint", "20", "主键", "否", "用户ID"),
        ("username", "varchar", "50", "", "否", "用户名"),
        ("password", "varchar", "100", "", "否", "密码"),
        ("nickname", "varchar", "50", "", "是", "昵称"),
        ("phone", "varchar", "20", "", "是", "手机号"),
        ("email", "varchar", "50", "", "是", "邮箱"),
        ("role_id", "bigint", "20", "外键", "否", "角色ID"),
        ("status", "tinyint", "4", "", "否", "状态(1启用/0停用)"),
        ("create_time", "datetime", "", "", "否", "创建时间"),
    )),
    ("表4-3 角色表（sys_role）结构", cols(
        ("role_id", "bigint", "20", "主键", "否", "角色ID"),
        ("role_name", "varchar", "50", "", "否", "角色名称"),
        ("role_desc", "varchar", "200", "", "是", "角色描述"),
    )),
    ("表4-4 商品分类表（category）结构", cols(
        ("category_id", "bigint", "20", "主键", "否", "分类ID"),
        ("category_name", "varchar", "50", "", "否", "分类名称"),
        ("parent_id", "bigint", "20", "", "是", "父分类ID"),
        ("sort", "int", "11", "", "是", "排序"),
    )),
    ("表4-5 商品表（goods）结构", cols(
        ("goods_id", "bigint", "20", "主键", "否", "商品ID"),
        ("goods_name", "varchar", "100", "", "否", "商品名称"),
        ("category_id", "bigint", "20", "外键", "否", "分类ID"),
        ("price", "decimal", "10,2", "", "否", "价格"),
        ("stock", "int", "11", "", "否", "库存"),
        ("image", "varchar", "200", "", "是", "图片"),
        ("description", "text", "", "", "是", "描述"),
        ("status", "tinyint", "4", "", "否", "状态(1上架/0下架)"),
        ("create_time", "datetime", "", "", "否", "创建时间"),
    )),
    ("表4-6 购物车表（cart）结构", cols(
        ("cart_id", "bigint", "20", "主键", "否", "购物车ID"),
        ("user_id", "bigint", "20", "", "否", "用户ID"),
        ("goods_id", "bigint", "20", "外键", "否", "商品ID"),
        ("quantity", "int", "11", "", "否", "数量"),
        ("create_time", "datetime", "", "", "否", "创建时间"),
    )),
    ("表4-7 订单表（orders）结构", cols(
        ("order_id", "bigint", "20", "主键", "否", "订单ID"),
        ("order_no", "varchar", "50", "", "否", "订单编号"),
        ("user_id", "bigint", "20", "", "否", "用户ID"),
        ("receiver_name", "varchar", "50", "", "否", "收货人"),
        ("receiver_phone", "varchar", "20", "", "否", "联系电话"),
        ("receiver_address", "varchar", "200", "", "否", "收货地址"),
        ("total_amount", "decimal", "10,2", "", "否", "订单金额"),
        ("status", "tinyint", "4", "", "否", "状态(0待支付/1待发货/2已发货/3已完成/4已取消)"),
        ("create_time", "datetime", "", "", "否", "下单时间"),
        ("pay_time", "datetime", "", "", "是", "支付时间"),
        ("ship_time", "datetime", "", "", "是", "发货时间"),
    )),
    ("表4-8 订单明细表（order_item）结构", cols(
        ("item_id", "bigint", "20", "主键", "否", "明细ID"),
        ("order_id", "bigint", "20", "外键", "否", "订单ID"),
        ("goods_id", "bigint", "20", "外键", "否", "商品ID"),
        ("goods_name", "varchar", "100", "", "否", "商品名称"),
        ("price", "decimal", "10,2", "", "否", "单价"),
        ("quantity", "int", "11", "", "否", "数量"),
        ("subtotal", "decimal", "10,2", "", "否", "小计"),
    )),
    ("表4-9 评价表（review）结构", cols(
        ("review_id", "bigint", "20", "主键", "否", "评价ID"),
        ("order_id", "bigint", "20", "外键", "否", "订单ID"),
        ("goods_id", "bigint", "20", "外键", "否", "商品ID"),
        ("user_id", "bigint", "20", "", "否", "用户ID"),
        ("rating", "tinyint", "4", "", "否", "评分(1-5)"),
        ("content", "varchar", "500", "", "是", "评价内容"),
        ("create_time", "datetime", "", "", "否", "评价时间"),
    )),
    ("表4-10 售后表（after_sale）结构", cols(
        ("sale_id", "bigint", "20", "主键", "否", "售后ID"),
        ("order_id", "bigint", "20", "外键", "否", "订单ID"),
        ("goods_id", "bigint", "20", "外键", "否", "商品ID"),
        ("user_id", "bigint", "20", "", "否", "用户ID"),
        ("type", "tinyint", "4", "", "否", "售后类型(1退货/2退款)"),
        ("reason", "varchar", "200", "", "否", "售后原因"),
        ("description", "varchar", "500", "", "是", "售后说明"),
        ("status", "tinyint", "4", "", "否", "状态(0待处理/1已同意/2已拒绝/3已完成)"),
        ("apply_time", "datetime", "", "", "否", "申请时间"),
        ("handle_time", "datetime", "", "", "是", "处理时间"),
    )),
    ("表4-11 物流表（shipping）结构", cols(
        ("shipping_id", "bigint", "20", "主键", "否", "物流ID"),
        ("order_id", "bigint", "20", "外键", "否", "订单ID"),
        ("tracking_no", "varchar", "50", "", "否", "物流单号"),
        ("company", "varchar", "50", "", "否", "物流公司"),
        ("ship_time", "datetime", "", "", "否", "发货时间"),
    )),
]

# ================= 工具 =================
def set_font(run, east="宋体", latin="Times New Roman", size=Pt(12), bold=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), east)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold

def find_para(text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise RuntimeError("未找到段落: " + text)

def insert_para_after(anchor_el, text="", style=None, align=None, size=Pt(12), bold=None):
    new_p = OxmlElement("w:p")
    anchor_el.addnext(new_p)
    p = Paragraph(new_p, body_parent)
    if style:
        p.style = doc.styles[style]
    if align is not None:
        p.alignment = align
    if text:
        for i, line in enumerate(text.split("\n")):
            if i > 0:
                r = p.add_run(); r.add_break()
            r = p.add_run(line)
            if style is None:
                set_font(r, size=size, bold=bold)
    return new_p

def add_heading2_after(anchor_el, text):
    new_p = OxmlElement("w:p")
    anchor_el.addnext(new_p)
    p = Paragraph(new_p, body_parent)
    p.style = doc.styles["Heading 2"]
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "1")
    numId = OxmlElement("w:numId"); numId.set(qn("w:val"), "4")
    numPr.append(ilvl); numPr.append(numId)
    pPr.append(numPr)
    p.add_run(text)
    return new_p

def insert_image_after(anchor_el, img_name, width_in=5.5, design=True):
    path = os.path.join(DIAG_D if design else DIAG, img_name)
    new_p = OxmlElement("w:p")
    anchor_el.addnext(new_p)
    p = Paragraph(new_p, body_parent)
    p.alignment = CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(width_in))
    return new_p

def create_table(data, col_widths=None, font_size=Pt(10.5)):
    rows = len(data); cols = len(data[0])
    tbl = doc.add_table(rows=rows, cols=cols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            _set_cell(cell, val, bold=(i == 0), size=font_size)
    if col_widths:
        for j, w in enumerate(col_widths):
            for i in range(rows):
                tbl.cell(i, j).width = Inches(w)
    return tbl

def _set_cell(cell, text, bold=False, size=Pt(10.5), align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    for i, line in enumerate(text.split("\n")):
        if i > 0:
            r = p.add_run(); r.add_break()
        r = p.add_run(line)
        set_font(r, size=size, bold=bold)

def move_table_after(tbl, anchor_el):
    anchor_el.addnext(tbl._tbl)

def caption_after(anchor_el, text):
    return insert_para_after(anchor_el, text, align=CENTER, size=Pt(10.5), bold=True)

# ================= 定位锚点 =================
p_aim = find_para("编写目的")
p_goal = find_para("项目目标")
p_req = find_para("主要软件需求")
p_arch = find_para("系统架构图")
p_biz = find_para("业务流程图")
p_class = find_para("类图")
p_module = find_para("模块描述")
p_er = find_para("E-R图")
p_logic = find_para("逻辑模型")
p_phys = find_para("数据表物理模型")

# 删除原模块描述空表
old_tbl = None
for t in doc.tables:
    if t.rows and t.rows[0].cells[0].text.strip() == "模块名称":
        old_tbl = t
        break
assert old_tbl is not None, "未找到模块描述表"
old_tbl._tbl.getparent().remove(old_tbl._tbl)

# ================= 1. 引言 =================
insert_para_after(p_aim._p, SEC11)
insert_para_after(p_goal._p, SEC12)

# ================= 2. 主要软件需求 =================
insert_para_after(p_req._p, SEC2)

# ================= 3.1 系统架构图 =================
cur = insert_para_after(p_arch._p, ARCH_TEXT)
cur = insert_image_after(cur, "01_arch.png", 5.8, design=True)
cur = caption_after(cur, "图3-1 系统架构图")

# ================= 3.2 业务流程图 =================
cur = insert_para_after(p_biz._p, BIZ_TEXT)
for img, cap in BIZ_FLOWS:
    cur = insert_image_after(cur, img, 2.3, design=False)
    cur = caption_after(cur, cap)

# ================= 3.3 类图 =================
cur = insert_para_after(p_class._p, CLASS_TEXT)
cur = insert_image_after(cur, "02_class.png", 4.6, design=True)
cur = caption_after(cur, "图3-6 系统类图")

# ================= 3.4 时序图（新增） =================
cur = add_heading2_after(cur, "时序图")
cur = insert_para_after(cur, SEQ_TEXT)
for img, cap in SEQ_DIAGRAMS:
    cur = insert_image_after(cur, img, 5.6, design=True)
    cur = caption_after(cur, cap)

# ================= 3.5 模块描述 =================
# 在模块描述标题下加表题与表格
cap = insert_para_after(p_module._p, "表3-1 系统模块描述", align=CENTER, size=Pt(10.5), bold=True)
mod_data = [MODULE_HEADER] + [list(m) for m in MODULES]
tbl = create_table(mod_data, col_widths=[1.0, 1.5, 2.8, 1.7], font_size=Pt(10.5))
move_table_after(tbl, cap)

# ================= 3.6 界面设计（新增） =================
cur = add_heading2_after(tbl._tbl, "界面设计")
cur = insert_para_after(cur, UI_TEXT)
for img, cap in UI_PAGES:
    w = 5.6 if "admin" in img else 5.0
    cur = insert_image_after(cur, img, w, design=True)
    cur = caption_after(cur, cap)

# ================= 4.1 E-R图 =================
cur = insert_para_after(p_er._p, ER_TEXT)
for img, cap in ER_DIAGRAMS:
    w = 3.0 if "permission" in img else 5.8
    cur = insert_image_after(cur, img, w, design=True)
    cur = caption_after(cur, cap)

# ================= 4.2 逻辑模型 =================
cur = insert_para_after(p_logic._p, LOGIC_TEXT)
cap = caption_after(cur, "表4-1 数据库表清单")
logic_data = [LOGIC_HEADER] + [list(r) for r in LOGIC_TABLES]
tbl_logic = create_table(logic_data, col_widths=[1.3, 1.3, 1.8, 2.6], font_size=Pt(10.5))
move_table_after(tbl_logic, cap)

# ================= 4.3 数据表物理模型 =================
cur = insert_para_after(p_phys._p, "各数据表的物理结构如表4-2至表4-11所示。")
for cap_text, rows in PHYSICAL:
    cap = caption_after(cur, cap_text)
    data = [PHYSICAL_HEADER] + rows
    t = create_table(data, col_widths=[1.2, 0.9, 0.7, 0.9, 0.8, 2.0], font_size=Pt(10.5))
    move_table_after(t, cap)
    cur = t._tbl

try:
    doc.save(OUT)
    print("saved:", OUT)
except PermissionError:
    OUT_TMP = OUT.replace(".docx", "-临时.docx")
    doc.save(OUT_TMP)
    print("原文件被占用，已保存到临时文件:", OUT_TMP)
