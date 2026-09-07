# -*- coding: utf-8 -*-
import os
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from lxml import etree

BASE = r"D:\20260907实训\商城系统\goods"

# (level, text, page)
REQ_TOC = [
    (1, "0. 文档介绍", 4),
    (2, "0.1 文档目的", 4),
    (2, "0.2 文档范围", 4),
    (2, "0.3 读者对象", 4),
    (2, "0.4 参考文档", 5),
    (2, "0.5 术语与缩写解释", 5),
    (1, "1. 产品介绍", 5),
    (1, "2. 产品面向的用户群体", 6),
    (1, "3. 产品应当遵循的标准或规范", 6),
    (1, "4. 产品的功能性需求", 7),
    (2, "4.1 产品业务流程图", 7),
    (2, "4.2 功能性需求分类", 9),
    (2, "4.3 用例描述", 11),
    (1, "5. 产品的非功能性需求", 14),
    (2, "5.1 用户界面需求", 14),
    (2, "5.2 软硬件环境需求", 15),
    (2, "5.3 产品质量需求", 15),
    (1, "6. 其他需求", 16),
]

DESIGN_TOC = [
    (1, "1. 引言", 2),
    (2, "1.1 编写目的", 2),
    (2, "1.2 项目目标", 2),
    (1, "2. 主要软件需求", 3),
    (1, "3. 软件系统结构设计", 3),
    (2, "3.1 系统架构图", 3),
    (2, "3.2 业务流程图", 4),
    (2, "3.3 类图", 5),
    (2, "3.4 时序图", 6),
    (2, "3.5 模块描述", 8),
    (2, "3.6 界面设计", 9),
    (1, "4. 数据库设计", 11),
    (2, "4.1 E-R图", 11),
    (2, "4.2 逻辑模型", 12),
    (2, "4.3 数据表物理模型", 13),
]


def make_toc_para(doc, level, text, page):
    new_p = OxmlElement('w:p')
    p = Paragraph(new_p, doc.paragraphs[0]._parent)
    p.style = doc.styles['toc %d' % level]
    # 文本
    run = p.add_run(text)
    set_font(run, size=docx.shared.Pt(10.5))
    # 制表符
    r2 = p.add_run()
    r2._r.append(OxmlElement('w:tab'))
    # 页码
    r3 = p.add_run(str(page))
    set_font(r3, size=docx.shared.Pt(10.5))
    return new_p


def set_font(run, east='宋体', latin='Times New Roman', size=None, bold=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn('w:eastAsia'), east)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold


def rebuild_toc(path, entries):
    doc = docx.Document(path)
    body = doc.element.body

    # 1. 找 toc 样式段落（缓存目录项）
    toc_paras = [p for p in doc.paragraphs if p.style.name.startswith('toc')]

    new_elems = [make_toc_para(doc, lv, tx, pg) for (lv, tx, pg) in entries]

    if toc_paras:
        # 有缓存项：在第一条前按正序插入新项，再删除旧项
        anchor = toc_paras[0]._p
        for e in new_elems:
            anchor.addprevious(e)
        for p in toc_paras:
            p._p.getparent().remove(p._p)
    else:
        # 无缓存项（设计文档：TOC 在 sdt 内容控件内，空）
        sdt = None
        for it in body.findall('.//' + qn('w:instrText')):
            if 'TOC' in (it.text or ''):
                # 向上找 sdt
                anc = it
                while anc is not None:
                    if anc.tag == qn('w:sdt'):
                        sdt = anc
                        break
                    anc = anc.getparent()
                break
        assert sdt is not None, '未找到 TOC 的 sdt'
        for e in new_elems:
            sdt.addprevious(e)
        sdt.getparent().remove(sdt)

    return doc


def save(path, doc):
    try:
        doc.save(path)
        print('已保存:', path)
    except PermissionError:
        tmp = path.replace('.docx', '-临时.docx')
        doc.save(tmp)
        print('原文件被占用，已存到临时文件:', tmp)


import docx.shared
for path, entries in [
    (os.path.join(BASE, 'docs', '需求分析规格说明书-完成版.docx'), REQ_TOC),
    (os.path.join(BASE, 'docs', '系统设计文档-完成版.docx'), DESIGN_TOC),
]:
    doc = rebuild_toc(path, entries)
    save(path, doc)
