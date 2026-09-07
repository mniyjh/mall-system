# -*- coding: utf-8 -*-
import os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

BASE = r"D:\20260907实训\商城系统\goods"
SRC = os.path.join(BASE, "docs", "需求分析规格说明书.docx")
OUT = os.path.join(BASE, "docs", "需求分析规格说明书-完成版.docx")
OUT_TMP = os.path.join(BASE, "docs", "需求分析规格说明书-完成版-临时.docx")
DIAG = os.path.join(BASE, "diagrams")

doc = docx.Document(SRC)
body_parent = doc.paragraphs[0]._parent

CENTER = WD_ALIGN_PARAGRAPH.CENTER

# ================= 内容数据 =================
SECTION1 = ("商城系统是一款基于浏览器/服务器（B/S）架构的网上商品交易系统。系统默认仅存在一家商户，"
            "由商户侧的运营人员负责商品的发布以及订单、发货、售后等业务的运营管理，管理员负责用户管理、"
            "角色管理等系统级内容的维护，买家则可在系统中完成商品的选购、下单、收货、评价与售后等操作。"
            "系统围绕商品上架、购买、发货、收货评价四大核心业务流程展开，为交易双方提供一个完整的在线商品交易平台。")

SECTION2_LEAD = "本系统默认仅有一家商户，共包含三类用户角色，各角色的主要职责如下："
SECTION2_ROLES = [
    "（1）管理员：系统的最高权限角色，负责系统级内容的维护与管理，主要承担用户管理、角色管理等职能，保障系统用户体系与权限体系的正常运转。",
    "（2）运营人员：商户侧的运营角色，负责商城日常业务的运营与管理，主要承担商品管理、订单管理、发货管理、售后管理等职能，保障商品的正常上架、交易与售后服务的顺利进行。",
    "（3）买家：商城的使用者，主要承担选购商品、购物车管理、下单、收货、评价、申请售后等职能，完成从选购商品到售后处理的一整套消费流程。",
]

PROCESSES = [
    ("商品上架流程：运营人员录入并完善商品信息，经审核无误后上架商品，商品进入可售状态。",
     "02_activity_shangjia.png", "图4-1 商品上架流程活动图"),
    ("购买流程：买家查询并选择商品，将其加入购物车，确认购买信息后提交订单并完成支付。",
     "03_activity_goumai.png", "图4-2 购买流程活动图"),
    ("发货流程：运营人员核对订单与商品信息，打包商品并登记物流单号，订单进入已发货状态。",
     "04_activity_fahuo.png", "图4-3 发货流程活动图"),
    ("收货评价流程：买家收到商品后确认收货并评价，如不满意可申请售后，由运营人员处理售后直至结束。",
     "05_activity_shouhuo.png", "图4-4 收货评价流程活动图"),
]

USECASE_DIAGRAMS = [
    ("图4-5 管理员用例图", "01a_use_case_admin.png"),
    ("图4-6 运营人员用例图", "01b_use_case_operator.png"),
    ("图4-7 买家用例图", "01c_use_case_buyer.png"),
]

USECASE_LIST_HEADER = ["序号", "用例名称", "用例功能描述", "备注"]
USECASE_LIST = [
    ("1", "用户管理", "维护系统用户账号，支持用户的查询、新增、修改、删除及停用/启用", "管理员"),
    ("2", "角色管理", "维护系统角色及其权限，支持角色的查询、新增、修改、删除及权限分配", "管理员"),
    ("3", "添加商品", "运营人员录入新商品的基本信息并保存到系统中", "运营人员"),
    ("4", "修改商品", "运营人员对已录入商品的信息进行修改", "运营人员"),
    ("5", "上架商品", "运营人员将审核通过的商品设为可售状态", "运营人员"),
    ("6", "下架商品", "运营人员将不再销售的商品设为停售状态", "运营人员"),
    ("7", "查询订单", "运营人员查询、浏览商城订单列表及订单详情", "运营人员"),
    ("8", "发货", "运营人员对已支付订单安排发货并登记物流信息", "运营人员"),
    ("9", "处理售后", "运营人员审核买家的售后申请并作出处理", "运营人员"),
    ("10", "查询商品", "买家浏览、检索商城商品列表及商品详情", "买家"),
    ("11", "添加到购物车", "买家将选中的商品加入购物车", "买家"),
    ("12", "下单", "买家确认购买信息并提交订单完成支付", "买家"),
    ("13", "确认收货", "买家收到商品后确认收货", "买家"),
    ("14", "评价商品", "买家对已收货的商品进行评价", "买家"),
    ("15", "申请售后", "买家对订单中的商品发起售后申请", "买家"),
]

UC_DESC = [
    {
        "caption": "表4-2 添加商品用例描述",
        "fields": [
            ("用例编号", "UC-001"), ("用例名称", "添加商品"), ("参与者", "运营人员"),
            ("用例说明", "运营人员录入新商品的基本信息，经校验通过后，商品进入待上架状态，供后续上架操作使用。"),
            ("前置条件", "运营人员已登录系统，且具备商品管理权限。"),
            ("基本事件流", "1. 运营人员发起添加商品请求。\n2. 系统展示商品信息录入要求。\n3. 运营人员录入商品的名称、类别、价格、库存、图片、描述等基本信息。\n4. 运营人员提交商品信息。\n5. 系统校验所录入信息的完整性与合法性。\n6. 商品信息校验通过并录入完成，商品进入待上架状态。\n7. 用例结束。"),
            ("替代事件流", "4a. 商品信息不完整或不合法：系统提示缺失或错误的信息，运营人员补全或更正后重新提交，返回基本事件流第5步。\n4b. 运营人员放弃添加：运营人员取消操作，用例结束，商品信息不作任何处理。"),
            ("后置条件", "商品信息保存成功，商品处于待上架状态。"),
        ],
    },
    {
        "caption": "表4-3 下单用例描述",
        "fields": [
            ("用例编号", "UC-002"), ("用例名称", "下单"), ("参与者", "买家"),
            ("用例说明", "买家确认购买信息后提交订单并完成支付，系统生成待发货订单。"),
            ("前置条件", "买家已登录系统，且购物车中存在待结算的商品。"),
            ("基本事件流", "1. 买家发起下单请求。\n2. 系统展示待结算商品、收货地址与应付金额。\n3. 买家确认购买信息并提交订单。\n4. 系统校验商品库存与订单信息。\n5. 买家完成支付。\n6. 系统确认支付成功，订单生成并进入待发货状态。\n7. 用例结束。"),
            ("替代事件流", "4a. 商品库存不足：系统提示库存不足信息，买家调整购买数量或放弃购买，返回基本事件流第3步或结束用例。\n5a. 支付失败：系统提示支付失败，买家可重新支付或取消订单，返回基本事件流第5步或结束用例。"),
            ("后置条件", "订单生成成功，订单处于待发货状态。"),
        ],
    },
    {
        "caption": "表4-4 发货用例描述",
        "fields": [
            ("用例编号", "UC-003"), ("用例名称", "发货"), ("参与者", "运营人员"),
            ("用例说明", "运营人员对已支付的待发货订单进行发货处理，登记物流信息，订单进入已发货状态。"),
            ("前置条件", "运营人员已登录系统，且存在已支付、待发货的订单。"),
            ("基本事件流", "1. 运营人员发起发货请求。\n2. 系统展示待发货订单及商品信息。\n3. 运营人员核对订单与商品信息。\n4. 运营人员打包商品并选择物流渠道。\n5. 运营人员登记物流单号。\n6. 物流信息登记完成，订单进入已发货状态。\n7. 用例结束。"),
            ("替代事件流", "3a. 订单信息与商品不符：运营人员中止发货操作，用例结束。\n5a. 物流单号不合法：系统提示物流单号有误，运营人员更正后重新登记，返回基本事件流第5步。"),
            ("后置条件", "订单物流信息登记完成，订单处于已发货状态。"),
        ],
    },
    {
        "caption": "表4-5 申请售后用例描述",
        "fields": [
            ("用例编号", "UC-004"), ("用例名称", "申请售后"), ("参与者", "买家"),
            ("用例说明", "买家对已收货的商品发起售后申请，说明售后原因与诉求，提交后等待运营人员处理。"),
            ("前置条件", "买家已登录系统，且存在已收货的订单商品。"),
            ("基本事件流", "1. 买家发起售后申请请求。\n2. 系统展示可申请售后的订单商品。\n3. 买家选择售后商品、售后类型与原因，填写售后说明。\n4. 买家提交售后申请。\n5. 系统校验售后申请的完整性。\n6. 售后申请提交完成，售后申请进入待处理状态。\n7. 用例结束。"),
            ("替代事件流", "4a. 买家放弃申请：买家取消操作，用例结束。\n5a. 售后申请信息不完整：系统提示缺失信息，买家补全后重新提交，返回基本事件流第5步。"),
            ("后置条件", "售后申请提交成功，售后申请处于待处理状态。"),
        ],
    },
]

DOC_PURPOSE = ("本文档旨在明确商城系统的功能需求与非功能需求，界定系统的用户角色、用例及核心业务流程，"
               "为后续的系统设计、开发与测试提供依据，同时作为项目相关方沟通与评审的统一标准。")

DOC_SCOPE = ("本文档描述了商城系统的需求，包括产品介绍、面向的用户群体、应遵循的标准与规范、功能性需求"
             "（四大业务流程、用例清单、用例图、用例描述）以及非功能性需求。本文档适用于商城系统的需求分析"
             "阶段，涵盖系统默认单一商户场景下的管理员、运营人员、买家三类角色，不涉及系统的具体实现技术与部署方案。")

UI_REQS = [
    ("界面风格", "整体界面简洁、美观，风格统一，符合主流电商网站的设计习惯"),
    ("页面布局", "前台采用商品展示、分类导航与搜索布局；后台管理采用左侧菜单与内容区布局"),
    ("导航设计", "导航层级清晰，用户可在三步内到达目标功能页面"),
    ("操作提示", "关键操作均有明确的结果提示，操作失败时给出具体原因"),
    ("输入校验", "所有输入项均有格式校验，必填项有醒目标识"),
    ("列表分页", "商品、订单等列表数据支持分页显示，每页显示数量合理"),
]

TERMS = [
    ("SRS", "Software Requirements Specification，软件需求规格说明书，本文档即属于此类文档"),
    ("B/S", "Browser/Server，浏览器/服务器架构，用户通过浏览器访问并操作系统"),
    ("用例（Use Case）", "描述参与者与系统之间一次交互的功能单元"),
    ("活动图（Activity Diagram）", "描述业务流程中各项活动及其流转关系的图形化表示"),
    ("参与者（Actor）", "与系统发生交互的外部角色，本系统中指管理员、运营人员、买家"),
    ("商城系统", "本文档所描述的网上商品交易系统，默认仅存在一家商户"),
]

# ================= 工具函数 =================
def set_font(run, east="宋体", latin="Times New Roman", size=Pt(12), bold=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), east)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold


def find_para(text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise RuntimeError("未找到段落: " + text)


def insert_para_after(anchor_el, text="", style=None, align=None, size=Pt(12), bold=None):
    new_p = OxmlElement("w:p")
    anchor_el.addnext(new_p)
    p = Paragraph(new_p, body_parent)
    if style:
        p.style = doc.styles[style]
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        if style is None:
            set_font(run, size=size, bold=bold)
    return new_p  # 返回 w:p 元素，便于链式插入


def insert_image_after(anchor_el, img_name, width_in=2.5):
    new_p = OxmlElement("w:p")
    anchor_el.addnext(new_p)
    p = Paragraph(new_p, body_parent)
    p.alignment = CENTER
    run = p.add_run()
    run.add_picture(os.path.join(DIAG, img_name), width=Inches(width_in))
    return new_p


def create_table(data, col_widths=None, font_size=Pt(10.5), bold_first_col=False):
    rows = len(data)
    cols = len(data[0])
    tbl = doc.add_table(rows=rows, cols=cols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            b = (i == 0) or (bold_first_col and j == 0)
            _set_cell(cell, val, bold=b, size=font_size)
    if col_widths:
        for j, w in enumerate(col_widths):
            for i in range(rows):
                tbl.cell(i, j).width = Inches(w)
    return tbl


def _set_cell(cell, text, bold=False, size=Pt(10.5), align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    for i, line in enumerate(text.split("\n")):
        if i > 0:
            r = p.add_run()
            r.add_break()
        r = p.add_run(line)
        set_font(r, size=size, bold=bold)


def move_table_after(tbl, anchor_el):
    anchor_el.addnext(tbl._tbl)


# ================= 定位锚点 =================
p41 = find_para("1. 产品介绍")
p42 = find_para("产品面向的用户群体")
p48 = find_para("以下是四大核心功能的流程图：")
p51 = find_para("4.2 功能性需求分类")

# 删除原有空表（功能类别|功能名称|描述）
old_tbl = None
for t in doc.tables:
    if t.rows and "功能类别" in t.rows[0].cells[0].text:
        old_tbl = t
        break
assert old_tbl is not None, "未找到用例清单空表"
old_tbl._tbl.getparent().remove(old_tbl._tbl)

# ================= 1. 产品介绍 =================
insert_para_after(p41._p, SECTION1)

# ================= 2. 角色说明 =================
cur = p42._p
cur = insert_para_after(cur, SECTION2_LEAD)
for role in SECTION2_ROLES:
    cur = insert_para_after(cur, role)

# ================= 4.1 活动图 =================
cur = p48._p
for desc, img, caption in PROCESSES:
    cur = insert_para_after(cur, desc)
    cur = insert_image_after(cur, img, width_in=2.5)
    cur = insert_para_after(cur, caption, align=CENTER, size=Pt(10.5), bold=True)

# ================= 4.2 用例图 + 用例清单 =================
cur = p51._p
cur = insert_para_after(cur, "本系统的用例图如图4-5至图4-7所示，据此得到系统的用例清单，详见表4-1。")
for caption, img in USECASE_DIAGRAMS:
    cur = insert_image_after(cur, img, width_in=3.2)
    cur = insert_para_after(cur, caption, align=CENTER, size=Pt(10.5), bold=True)

# 用例清单表题 + 表格
cap = insert_para_after(cur, "表4-1 系统用例清单", align=CENTER, size=Pt(10.5), bold=True)
list_data = [USECASE_LIST_HEADER] + [list(row) for row in USECASE_LIST]
tbl_list = create_table(list_data, col_widths=[0.6, 1.5, 3.1, 0.9], font_size=Pt(10.5))
move_table_after(tbl_list, cap)

# ================= 4.3 用例描述 =================
cur = insert_para_after(tbl_list._tbl, "4.3 用例描述", style="Heading 2")
cur = insert_para_after(cur, "本系统共包含15个用例，现就四大业务流程中的添加商品、下单、发货、申请售后四个核心用例描述如下。")
for desc in UC_DESC:
    cap = insert_para_after(cur, desc["caption"], align=CENTER, size=Pt(10.5), bold=True)
    data = [["名称", "描述"]] + [[k, v] for k, v in desc["fields"]]
    tbl = create_table(data, col_widths=[1.3, 4.7], font_size=Pt(10.5), bold_first_col=True)
    move_table_after(tbl, cap)
    cur = tbl._tbl

# ================= 0.1 文档目的 / 0.2 文档范围 =================
p_01 = find_para("0.1 文档目的")
p_02 = find_para("0.2 文档范围")
insert_para_after(p_01._p, DOC_PURPOSE)
insert_para_after(p_02._p, DOC_SCOPE)

# ================= 5.1 用户界面需求 =================
ui_tbl = None
for t in doc.tables:
    if (t.rows and t.rows[0].cells[0].text.strip() == "需求名称"
            and len(t.rows) > 1 and t.rows[1].cells[0].text.strip() == ""):
        ui_tbl = t
        break
assert ui_tbl is not None, "未找到用户界面需求空表"
for i, (name, detail) in enumerate(UI_REQS):
    row = ui_tbl.rows[i + 1]
    _set_cell(row.cells[0], name)
    _set_cell(row.cells[1], detail)

# ================= 0.5 术语与缩写解释 =================
term_tbl = None
for t in doc.tables:
    if t.rows and "缩写" in t.rows[0].cells[0].text:
        term_tbl = t
        break
assert term_tbl is not None, "未找到术语表"
have = len(term_tbl.rows) - 1  # 去掉表头
for _ in range(len(TERMS) - have):
    term_tbl.add_row()
for i, (abbr, desc) in enumerate(TERMS):
    row = term_tbl.rows[i + 1]
    _set_cell(row.cells[0], abbr)
    _set_cell(row.cells[1], desc)

try:
    doc.save(OUT)
    print("saved:", OUT)
except PermissionError:
    doc.save(OUT_TMP)
    print("原文件被占用，已保存到临时文件:", OUT_TMP)
